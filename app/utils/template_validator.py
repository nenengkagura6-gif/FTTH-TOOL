"""
Validasi template Excel yang diunggah pengguna.

Engine BOQ dan HPDB menulis ke alamat sel yang tetap (mis. C2, O5, N6).
Template dengan susunan lain dulu diterima begitu saja dan angkanya
mendarat di sel yang salah tanpa pesan apa pun. Pemeriksaan di sini
menolak template yang jelas tidak cocok sebelum diproses.
"""
import io
import zipfile
from typing import List

from openpyxl import load_workbook

BOQ_REQUIRED_SHEETS = ("BoM AE", "BoQ NRO Cluster")


class TemplateError(ValueError):
    """Template tidak bisa dipakai; pesannya layak ditampilkan ke user."""


def _open(template_bytes: bytes):
    if not template_bytes:
        raise TemplateError("Template kosong.")
    if not zipfile.is_zipfile(io.BytesIO(template_bytes)):
        raise TemplateError(
            "Template harus berformat .xlsx. Format .xls (Excel 97-2003) tidak "
            "didukung — buka di Excel lalu simpan ulang sebagai .xlsx."
        )
    try:
        return load_workbook(io.BytesIO(template_bytes))
    except Exception as exc:  # noqa: BLE001
        raise TemplateError(f"Template tidak bisa dibuka: {exc}")


def validate_boq_template(template_bytes: bytes) -> None:
    wb = _open(template_bytes)
    missing: List[str] = [s for s in BOQ_REQUIRED_SHEETS if s not in wb.sheetnames]
    if missing:
        raise TemplateError(
            "Template BOQ tidak cocok: sheet "
            + ", ".join(f"'{s}'" for s in missing)
            + " tidak ditemukan. Sheet yang ada: "
            + ", ".join(wb.sheetnames)
        )


def validate_hpdb_template(template_bytes: bytes) -> None:
    wb = _open(template_bytes)
    ws = wb.worksheets[0]
    # Data HPDB mulai di baris 10; template yang lebih pendek dari itu
    # hampir pasti bukan template HPDB.
    if ws.max_row < 9:
        raise TemplateError(
            f"Template HPDB tidak cocok: sheet pertama '{ws.title}' hanya "
            f"{ws.max_row} baris, sedangkan data ditulis mulai baris 10."
        )


def validate_template(path) -> bool:
    """Kompatibilitas lama: validasi template BOQ dari path file."""
    with open(path, "rb") as fh:
        validate_boq_template(fh.read())
    return True
