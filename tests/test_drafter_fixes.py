"""Bukti perbaikan hasil audit tool Drafter (2026-09-25).

Jalankan dari akar repo:  python tests/test_drafter_fixes.py
Tidak memakai jaringan: geocoding dipalsukan.
"""
import io
import os
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "app"))

from openpyxl import load_workbook  # noqa: E402

failures = 0


def check(name, cond, detail=""):
    global failures
    print(("  [OK]    " if cond else "  [GAGAL] ") + f"{name} {detail}")
    if not cond:
        failures += 1


def kml(*parts: str, ns: bool = True) -> bytes:
    body = "".join(parts)
    xmlns = ' xmlns="http://www.opengis.net/kml/2.2"' if ns else ""
    return f'<?xml version="1.0" encoding="UTF-8"?><kml{xmlns}><Document>{body}</Document></kml>'.encode()


def pt(name, lon, lat):
    return f"<Placemark><name>{name}</name><Point><coordinates>{lon},{lat},0</coordinates></Point></Placemark>"


def line(name, coords):
    c = " ".join(f"{lo},{la},0" for lo, la in coords)
    return f"<Placemark><name>{name}</name><LineString><coordinates>{c}</coordinates></LineString></Placemark>"


def folder(name, *children):
    return f"<Folder><name>{name}</name>{''.join(children)}</Folder>"


# =====================================================================
print("=== 1. Pemuat bersama: zip bomb KMZ & DOCTYPE ===")
from utils.commons import load_kml_text, KmlLoadError  # noqa: E402

buf = io.BytesIO()
with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("doc.kml", b"<kml>" + b" " * (5 * 1024 * 1024) + b"</kml>")
try:
    load_kml_text(buf.getvalue(), True)
    check("KMZ dengan rasio kompresi ekstrem ditolak", False)
except KmlLoadError as e:
    check("KMZ dengan rasio kompresi ekstrem ditolak", True, f"-> {e}")

xxe = ('<?xml version="1.0"?><!DOCTYPE kml [<!ENTITY x SYSTEM "file:///etc/passwd">]>'
       '<kml><Document><name>&x;</name></Document></kml>').encode()
teks = load_kml_text(xxe)
check("DOCTYPE/entitas eksternal dibuang", "DOCTYPE" not in teks and "ENTITY" not in teks)

# =====================================================================
print("\n=== 2. KML to BOQ ===")
from engines.kml_engine import process_kml_to_excel  # noqa: E402

boq_kml = kml(
    folder("FDT 01",
           folder("LINE A",
                  folder("DISTRIBUTION CABLE",
                         line("LINE A 24C seg1", [(106.8, -6.2), (106.8001, -6.2)]),
                         line("LINE A 24C seg2", [(106.8001, -6.2), (106.8002, -6.2)]),
                         line("LINE A 96C", [(106.8, -6.2), (106.801, -6.2)])),
                  folder("FAT", pt("A01", 106.8, -6.2), pt("A02", 106.8001, -6.2)))),
    folder("FDT 04",
           folder("LINE A", folder("FAT", pt("A01", 106.9, -6.2)))),
)
res = process_kml_to_excel(boq_kml, "CBN01 UJI.kml")
wb = load_workbook(io.BytesIO(res["content"]))
ae = wb["BoM AE"]
warn = " | ".join(res["report"]["warnings"])
check("dua segmen ~11 m dijumlah dulu baru dibulatkan", ae["C2"].value == 22, f"-> C2={ae['C2'].value}")
check("kabel 96C dilaporkan, tidak hilang diam-diam", "96C" in warn)
check("FDT 04 dilaporkan di luar template", "FDT 04" in warn)
check("FAT FDT 01 tetap di kolom C", ae["C36"].value == 2, f"-> C36={ae['C36'].value}")

from utils.template_validator import validate_boq_template, TemplateError  # noqa: E402
try:
    validate_boq_template(b"\xd0\xcf\x11\xe0 bukan xlsx")
    check("template .xls ditolak dengan pesan jelas", False)
except TemplateError as e:
    check("template .xls ditolak dengan pesan jelas", ".xlsx" in str(e))

# =====================================================================
print("\n=== 3. KML to Database HP ===")
from engines import apd_engine  # noqa: E402

apd_engine.APDEngine.reverse_geocode = lambda self, lat, lon: {
    "province": "JAWA BARAT", "kabupaten": f"KAB{round(float(lon), 3)}", "kecamatan": "K",
    "desa": "D", "kodepos": "1", "jalan": f"JL {round(float(lat), 4)}"}

hp_kml = kml(
    folder("FDT", pt("FDT 01", 106.80, -6.20), pt("FDT 02", 106.90, -6.20)),
    folder("FDT 01",
           folder("LINE A",
                  folder("FAT", pt("FAT A01", 106.8000, -6.2000)),
                  folder("POLE", pt("P-FDT1", 106.80001, -6.2000)),
                  folder("HP COVER",
                         folder("A01", pt("HP 1", 106.8001, -6.2001), pt("HP 2", 106.8002, -6.2001)),
                         folder("A02", pt("HP 3", 106.8003, -6.2001)),
                         folder("A01", pt("HP 4", 106.8004, -6.2001))))),
    folder("FDT 02",
           folder("LINE A",
                  folder("FAT", pt("FAT A01", 106.9000, -6.2000)),
                  folder("POLE", pt("P-FDT2", 106.90001, -6.2000)),
                  folder("HP COVER", folder("A01", pt("HP X", 106.9001, -6.2001))))),
)
res = apd_engine.process_apd_hpdb(hp_kml, "APD_HPDB_CBN01 UJI.kml")
check("proses HPDB sukses", res["status"] == "success", res.get("message", ""))
wb = load_workbook(io.BytesIO(res["content"]))
sheets = {ws.title: ws for ws in wb.worksheets}
ws1 = sheets.get("Homepass Database FDT 01")
ws2 = sheets.get("Homepass Database FDT 02")
check("FDT dari folder kakek (FDT 02/LINE A/HP COVER) -> sheet FDT 02", ws2 is not None, str(list(sheets)))
if ws1 is not None and ws2 is not None:
    rows1 = [[c.value for c in r] for r in ws1.iter_rows(min_row=10) if r[6].value]
    ports_a01 = [r[7] for r in rows1 if r[6] == "A01"]
    check("HP A01 yang tidak berurutan: port 1,2,3 tanpa reset", ports_a01 == [1, 2, 3], f"-> {ports_a01}")
    fdt_ports = [r[1] for r in rows1 if r[1] not in (None, "")]
    check("port FDT tidak dobel", len(fdt_ports) == len(set(fdt_ports)), f"-> {fdt_ports}")
    check("A01 FDT 01 -> tiang FDT 01", rows1[0][8] == "P-FDT1", f"-> {rows1[0][8]}")
    row2 = [c.value for c in next(ws2.iter_rows(min_row=10, max_row=10))]
    check("A01 FDT 02 -> tiang FDT 02 (tidak tertimpa)", row2[8] == "P-FDT2", f"-> {row2[8]}")
    jalan = {ws1[f"AD{r}"].value for r in range(10, 10 + len(rows1))}
    check("geocoding per FAT (A01 & A02 beda jalan)", len(jalan) == 2, f"-> {jalan}")

# =====================================================================
print("\n=== 4. Pole Auto-Sorter ===")
from engines.pole_sorter_engine import (  # noqa: E402
    process_pole_sorter, _is_new_pole_folder, _is_existing_pole_folder)

check("'INPUT' bukan folder tiang baru", not _is_new_pole_folder("INPUT"))
check("'EXISTING CABLE' bukan folder tiang", not _is_existing_pole_folder("EXISTING CABLE"))
check("'EXT MR RW 04' folder tiang eksisting", _is_existing_pole_folder("EXT MR RW 04"))

ps_kml = kml(
    folder("FDT", pt("FDT 01", 106.8000, -6.2)),
    folder("LINE A FDT 01",
           folder("DISTRIBUTION CABLE", line("CABLE", [(106.8000, -6.2), (106.8030, -6.2)])),
           folder("NEW POLE", pt("P003", 106.8005, -6.2), pt("P001", 106.8025, -6.2),
                  pt("P002", 106.8015, -6.2)),
           folder("EXISTING CABLE", line("jangan diubah", [(106.8, -6.21), (106.81, -6.21)]))),
)
res = process_pole_sorter(ps_kml, "uji.kml")
with zipfile.ZipFile(io.BytesIO(res["content"])) as z:
    out = z.read("doc.kml").decode()
order = [n for n in ("P001", "P002", "P003")]
idx = {n: out.find(f"<name>{n}</name>") for n in order}
# Setelah urut jalur kabel: tiang terdekat FDT (106.8005) harus jadi P001.
pos_1 = out.find("106.8005")
check("tiang terdekat FDT jadi P001 (urut sepanjang kabel)",
      out.rfind("<name>P001</name>", 0, pos_1) > out.rfind("<name>P00", 0, out.rfind("<name>P001</name>", 0, pos_1)) - 1
      and out[:pos_1].rfind("<name>") == out[:pos_1].rfind("<name>P001</name>"))
check("folder EXISTING CABLE tidak disentuh", "jangan diubah" in out)
check("metode urut dilaporkan", res["report"]["groups"][0]["sort_method"] == "cable_route")

# =====================================================================
print("\n=== 5. Insert Coding ===")
from engines.insert_coding_engine import process_insert_coding  # noqa: E402

ic_kml = kml(
    folder("FDT", pt("FDT 1", 1, 1), pt("fdt 4", 1, 1), pt("FDT 5", 1, 1)),
    folder("LINE A FDT 04",
           folder("FAT", pt("FAT A02", 1, 1), pt("FAT A01", 1, 1)),
           folder("CABLE", line("24C", [(1, 1), (2, 2)]))),
)
res = process_insert_coding(ic_kml, "uji.kml", {1: "AAA.001", 4: "DDD.004"})
with zipfile.ZipFile(io.BytesIO(res["content"])) as z:
    out = z.read("doc.kml").decode()
check("prefix FDT 04 dipakai (bukan hanya 1-3)", "DDD.004.A01" in out and "DDD.004 - 24C" in out)
check("titik 'fdt 4' pakai prefix FDT 04", "<name>DDD.004</name>" in out)
check("FDT tanpa prefix dilaporkan", any("FDT 05" in w for w in res["report"]["warnings"]))

# =====================================================================
print("\n=== 6. KML-APD ===")
from engines.kml_apd_engine import process_kml_apd  # noqa: E402

poly = ("<Placemark><name>b</name><Polygon><outerBoundaryIs><LinearRing><coordinates>"
        "106.7999,-6.2001,0 106.8011,-6.2001,0 106.8011,-6.1999,0 106.7999,-6.1999,0 106.7999,-6.2001,0"
        "</coordinates></LinearRing></outerBoundaryIs></Polygon></Placemark>")
apd_kml = kml(folder(
    "PROYEK UJI",  # dibungkus satu folder proyek, seperti hasil Save Place As
    folder("FDT", pt("FDT 01", 106.8000, -6.2000)),
    folder("POLE", pt("T1", 106.8002, -6.2000), pt("T2", 106.8006, -6.2000), pt("T3", 106.8010, -6.2000)),
    folder("HP", pt("H1", 106.8005, -6.20005)),
    folder("LINE A FDT 01",
           folder("BOUNDARY", poly),
           folder("DISTRIBUTION CABLE", line("CABLE LINE A", [(106.8000, -6.2), (106.8010, -6.2)])),
           folder("SLING WIRE", line("SLING MANUAL", [(106.8002, -6.2), (106.8010, -6.2)]))),
))
res = process_kml_apd(apd_kml, "uji.kml")
check("KML berbungkus folder proyek tetap diproses", res["status"] == "success", res.get("message", ""))
if res["status"] == "success":
    out = res["content"].decode()
    check("xmlns KML dipertahankan di hasil", 'xmlns="http://www.opengis.net/kml/2.2"' in out)
    check("sling manual TIDAK diganti nama jadi CABLE", "SLING MANUAL" in out)
    check("FAT dibuat", res["report"]["stats"].get("fat") == 1, str(res["report"]["stats"]))
    check("slack FDT = 1 (bukan jumlah semua FDT)", "(1 slack FDT &amp; 1 slack FAT)" in out)

    # BasicMap harus bisa membaca hasil KML-APD
    from engines.kml2basicmap import Kml  # noqa: E402
    k = Kml(res["content"])
    check("BasicMap membaca hasil KML-APD", len(k.placemarks()) > 0, f"-> {len(k.placemarks())} placemark")

from engines.kml2basicmap import Kml  # noqa: E402
k = Kml(kml(folder("HP", pt("H1", 106.8, -6.2)), ns=False))
check("BasicMap membaca KML tanpa xmlns", len(k.placemarks()) == 1)

print("\nHASIL:", "SEMUA LULUS" if failures == 0 else f"{failures} KEGAGALAN")
sys.exit(1 if failures else 0)
